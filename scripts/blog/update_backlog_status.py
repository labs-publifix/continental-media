#!/usr/bin/env python3
"""
Continental Media — updates one row's Estado (and optionally Fecha de
publicación) in content/blog/Backlog_Blog_Continental_Media.xlsx.

Uses openpyxl WITHOUT data_only=True so formulas (Mes de publicación,
Día, Semana, Lote are all formula-derived columns — see the Esquema
sheet) and cell formatting survive the round-trip. This is the ONLY
script that writes to the backlog; scripts/lib/blog-backlog.mjs (Node)
is read-only and is what every build script uses to read it back.

Usage:
    python3 scripts/blog/update_backlog_status.py E-01 "En redacción"
    python3 scripts/blog/update_backlog_status.py E-01 Publicado --fecha 2026-10-19

Valid Estado values (per Esquema): Idea, Aprobado, En redacción, En revisión, Publicado.
"""
import argparse
import datetime
import sys
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parents[2]
BACKLOG_PATH = REPO_ROOT / "content" / "blog" / "Backlog_Blog_Continental_Media.xlsx"

VALID_ESTADOS = {"Idea", "Aprobado", "En redacción", "En revisión", "Publicado"}

# Column letters on the Backlog sheet (see Esquema sheet / blog-backlog.mjs).
COL_ID = "A"
COL_ESTADO = "W"
COL_FECHA = "X"


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("id", help="Backlog ID, e.g. E-01 or G-07")
    parser.add_argument("estado", choices=sorted(VALID_ESTADOS), help="New Estado value")
    parser.add_argument("--fecha", help="Override Fecha de publicación (YYYY-MM-DD). Only meaningful when estado=Publicado.")
    parser.add_argument("--path", default=str(BACKLOG_PATH), help="Path to the .xlsx (default: content/blog/Backlog_Blog_Continental_Media.xlsx)")
    args = parser.parse_args()

    path = Path(args.path)
    if not path.exists():
        print(f"ERROR: backlog not found at {path}", file=sys.stderr)
        sys.exit(1)

    wb = openpyxl.load_workbook(path)  # data_only defaults to False: formulas preserved
    if "Backlog" not in wb.sheetnames:
        print("ERROR: 'Backlog' sheet not found", file=sys.stderr)
        sys.exit(1)
    ws = wb["Backlog"]

    target_row = None
    for row in range(2, ws.max_row + 1):
        cell_id = ws[f"{COL_ID}{row}"].value
        if cell_id and str(cell_id).strip() == args.id:
            target_row = row
            break

    if target_row is None:
        print(f"ERROR: ID '{args.id}' not found in Backlog sheet", file=sys.stderr)
        sys.exit(1)

    old_estado = ws[f"{COL_ESTADO}{target_row}"].value
    ws[f"{COL_ESTADO}{target_row}"] = args.estado

    if args.fecha:
        try:
            y, m, d = (int(x) for x in args.fecha.split("-"))
            ws[f"{COL_FECHA}{target_row}"] = datetime.datetime(y, m, d)
        except ValueError:
            print(f"ERROR: --fecha must be YYYY-MM-DD, got '{args.fecha}'", file=sys.stderr)
            sys.exit(1)

    wb.save(path)
    fecha_note = f", fecha -> {args.fecha}" if args.fecha else ""
    print(f"OK: {args.id} Estado '{old_estado}' -> '{args.estado}'{fecha_note} (row {target_row})")


if __name__ == "__main__":
    main()
